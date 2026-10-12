"""Phase 6: 専門医レビュー用シートを出力する。

使い方（リポジトリ直下で）:
    .\\.venv\\Scripts\\python.exe scripts\\export_review.py ES          # review/review_ES_YYYYMMDD.xlsx
    .\\.venv\\Scripts\\python.exe scripts\\export_review.py ES --csv    # CSV（UTF-8 BOM付き）も出す

retired のカードは出力しない。記入後は scripts/import_review.py で取り込む。
"""
from __future__ import annotations

import csv
import datetime as dt
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
FACTS = ROOT / "data" / "facts"
REVIEW = ROOT / "review"

JUDGMENTS = ["OK", "修正", "削除"]
# (見出し, 列幅, 記入欄か)
COLUMNS = [
    ("ID", 9, False), ("status", 14, False), ("主張文", 60, False), ("条件", 26, False),
    ("鑑別・よくある誤解", 34, False), ("出典", 40, False), ("DOI", 22, False), ("PMID", 11, False),
    ("source_class", 12, False), ("uncertainty", 30, False), ("コアカリ", 26, False), ("出題基準", 26, False),
    ("Claude Codeによる修正", 40, False),
    ("判定", 8, True), ("修正案", 40, True), ("コメント", 30, True),
]
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")


def fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dict):
        return "；".join(f"{k}: {fmt(x)}" for k, x in v.items())
    if isinstance(v, list):
        return "；".join(fmt(x) for x in v)
    return str(v)


def claude_edits(card: dict) -> str:
    """Claude Code が主張文・出典を変えた履歴（専門医に特に見てほしい箇所）。"""
    keys = ("内容照合による修正", "起票：", "統合：本文を", "status draft → verified_sources")
    return "\n".join(f"[{h['date']}] {h['change']}" for h in card.get("history") or []
                     if h.get("by") == "claude-code" and any(k in h["change"] for k in keys))


def rows_for(prefix: str) -> list[list]:
    cards = yaml.safe_load((FACTS / f"{prefix}.yaml").read_text(encoding="utf-8"))
    rows = []
    for c in cards:
        if c["publication_info"]["status"] == "retired":
            continue
        a, p = c["assertion"], c["provenance"]
        rows.append([
            c["id"], c["publication_info"]["status"], a.get("plain_statement_ja"), fmt(a.get("condition")),
            "\n".join(x for x in (f"鑑別：{fmt(a.get('differential'))}" if a.get("differential") else "",
                                  f"誤解：{fmt(a.get('common_misconception'))}" if a.get("common_misconception") else "") if x),
            p.get("source"), p.get("doi") or "", p.get("pmid") or "", p.get("source_class"),
            fmt(c.get("uncertainty")), (c.get("curriculum") or {}).get("core_curriculum") or "",
            (c.get("curriculum") or {}).get("kokushi") or "", claude_edits(c), "", "", "",
        ])
    return rows


def write_xlsx(prefix: str, rows: list[list], path: Path) -> None:
    wb = Workbook()
    info = wb.active
    info.title = "説明"
    info_rows = [
        ("専門医レビューシート", f"領域 {prefix}　出力日 {dt.date.today().isoformat()}　{len(rows)}件"),
        ("確認者（氏名）", ""),
        ("確認日（YYYY-MM-DD）", ""),
        ("", ""),
        ("記入方法", "シート「review」の黄色の列（判定・修正案・コメント）に記入してください。"),
        ("判定 = OK", "内容を確認済み。取り込み時に verified_expert に昇格し、確認者・確認日を記録します。"),
        ("判定 = 修正", "「修正案」に正しい内容を書いてください。本文は自動では書き換えず、proposals/ に修正案として出します。"),
        ("判定 = 削除", "廃止の提案として proposals/ に出します（カードは削除せず、承認後に retired にします）。"),
        ("空欄", "未確認として扱い、何も変更しません。"),
        ("注意", "「Claude Codeによる修正」列に記載があるカードは、Claude Code が一次資料をもとに主張文や出典を変更・作成したものです。特に確認をお願いします。"),
        ("注意", "status が draft のカードは、出典との照合が未了です（理由は uncertainty 列）。"),
    ]
    for r in info_rows:
        info.append(list(r))
    for cell in ("B2", "B3"):
        info[cell].fill = INPUT_FILL
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 100
    for row in info.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    info["A1"].font = Font(bold=True, size=13)

    ws = wb.create_sheet("review")
    ws.append([c[0] for c in COLUMNS])
    for i, (name, width, is_input) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        cell.fill = INPUT_FILL if is_input else HEAD_FILL
        ws.column_dimensions[cell.column_letter].width = width
    for r in rows:
        ws.append(r)
    doi_col = [c[0] for c in COLUMNS].index("DOI") + 1
    pmid_col = [c[0] for c in COLUMNS].index("PMID") + 1
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if COLUMNS[cell.column - 1][2]:
                cell.fill = INPUT_FILL
        doi = row[doi_col - 1]
        if doi.value:
            doi.hyperlink = f"https://doi.org/{doi.value}"
            doi.font = Font(color="0563C1", underline="single")
        pmid = row[pmid_col - 1]
        if pmid.value:
            pmid.hyperlink = f"https://pubmed.ncbi.nlm.nih.gov/{pmid.value}/"
            pmid.font = Font(color="0563C1", underline="single")
    judge_col = ws.cell(row=1, column=[c[0] for c in COLUMNS].index("判定") + 1).column_letter
    dv = DataValidation(type="list", formula1=f'"{",".join(JUDGMENTS)}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{judge_col}2:{judge_col}{len(rows) + 1}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main(argv: list[str]) -> None:
    prefixes = [a for a in argv if not a.startswith("--")] or sorted(p.stem for p in FACTS.glob("*.yaml") if not p.stem.startswith("_"))
    stamp = dt.date.today().strftime("%Y%m%d")
    for prefix in prefixes:
        rows = rows_for(prefix)
        xlsx = REVIEW / f"review_{prefix}_{stamp}.xlsx"
        write_xlsx(prefix, rows, xlsx)
        print(f"{xlsx.relative_to(ROOT)}: {len(rows)} rows")
        if "--csv" in argv:
            path = xlsx.with_suffix(".csv")
            with path.open("w", encoding="utf-8-sig", newline="") as f:
                w = csv.writer(f)
                w.writerow([c[0] for c in COLUMNS])
                w.writerows(rows)
            print(f"{path.relative_to(ROOT)}: {len(rows)} rows")


if __name__ == "__main__":
    main(sys.argv[1:])
