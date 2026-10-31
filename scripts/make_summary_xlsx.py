"""事実票の取り込み結果を1つの Excel ファイルにまとめる（概要・出題基準の網羅・カード一覧・出典一覧・説明）。

使い方: python scripts/make_summary_xlsx.py   → reports/factcards_summary_<日付>.xlsx（件数は数式。LibreOffice で再計算すると値が入る）
"""
import yaml,glob,re,datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
R=Path(__file__).resolve().parent.parent
doms=[]
for l in (R/'data/domains.yaml').read_text(encoding='utf-8').splitlines():
    m=re.match(r'- \{prefix: (\w+), name: "?(.*?)"?, field: (.*?)\}',l)
    if m: doms.append(m.groups())
field_of={p:f for p,n,f in doms}
cards=[]
for p,n,f in doms:
    fp=R/'data/facts'/f'{p}.yaml'
    if fp.exists(): cards+=yaml.safe_load(fp.read_text(encoding='utf-8')) or []
S=yaml.safe_load((R/'data/sources/sources.yaml').read_text(encoding='utf-8'))
K=yaml.safe_load((R/'data/curriculum/kokushi_r6_items.yaml').read_text(encoding='utf-8'))
ORDER=['Ⅰ','Ⅱ','Ⅲ','Ⅳ','Ⅴ','Ⅵ','Ⅶ','Ⅷ','Ⅸ','Ⅹ','Ⅺ','Ⅻ','ⅩⅢ']
chs=sorted(K['chapters'],key=lambda c:ORDER.index(c['id']))
F=Font(name='Arial',size=10); FB=Font(name='Arial',size=10,bold=True); FH=Font(name='Arial',size=10,bold=True,color='FFFFFF')
HF=PatternFill('solid',fgColor='305496'); TF=PatternFill('solid',fgColor='D9E1F2')
thin=Side(style='thin',color='BFBFBF'); BD=Border(bottom=thin)
wb=Workbook()
def sheet(ws,headers,rows,widths,wrap=()):
    ws.append(headers)
    for c in ws[1]: c.font=FH; c.fill=HF; c.alignment=Alignment(vertical='center',wrap_text=True)
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font=F
            if c.column in wrap: c.alignment=Alignment(wrap_text=True,vertical='top')
            else: c.alignment=Alignment(vertical='top')
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
# --- カード一覧
wc=wb.active; wc.title='カード一覧'
rows=[]
for c in cards:
    a=c['assertion']; pv=c.get('provenance') or {}; cu=c.get('curriculum') or {}; pi=c.get('publication_info') or {}
    cond=a.get('condition') or {}
    rows.append([c['id'],field_of.get(c['id'][:2],''),c['id'][:2],c.get('subdomain') or '',a.get('plain_statement_ja',''),
        (cond.get('region') if isinstance(cond,dict) else '') or '',cu.get('kokushi') or '',cu.get('core_curriculum') or '',
        pv.get('source_ref') or '',pv.get('source') or '',pv.get('doi') or '',str(pv.get('pmid') or ''),pv.get('source_class') or '',
        pi.get('status') or '',pi.get('superseded_by') or '',str(pi.get('created') or ''),c.get('uncertainty') or ''])
sheet(wc,['ID','分野','接頭辞','主題（subdomain）','主張文','地域','出題基準（医学各論）','コアカリ','出典番号','出典','DOI','PMID','ライセンス区分','状態','置き換え先','作成日','不確かさ（uncertainty）'],
      rows,[9,14,6,18,70,8,40,30,10,40,22,10,12,12,10,11,50],wrap=(5,7,17))
NC=len(rows)+1
# --- 出典一覧
ws=wb.create_sheet('出典一覧')
srows=[[s['id'],s.get('title') or s.get('citation') or '',s.get('doi') or '',str(s.get('pmid') or ''),s.get('source_class') or '',s.get('license_note') or '',
        f'=COUNTIF(カード一覧!$I$2:$I${NC},A{i+2})'] for i,s in enumerate(S)]
sheet(ws,['出典番号','書誌','DOI','PMID','ライセンス区分','ライセンス表示','使っているカード数'],srows,[11,80,26,11,13,30,12],wrap=(2,))
NS=len(srows)+1
# --- 出題基準の網羅
wk=wb.create_sheet('出題基準の網羅')
krows=[]
for ch in chs:
    for g in ch['groups']:
        for it in g['items']:
            i=len(krows)+2
            krows.append([ch['id'],ch['name'],g['name'],it['code'],it['name'],it['level'],
                f'=COUNTIFS(カード一覧!$G$2:$G${NC},"*医学各論 "&D{i}&"（*",カード一覧!$N$2:$N${NC},"<>retired")',
                f'=IF(G{i}>0,"あり","なし")'])
sheet(wk,['章','章の名前','大項目','番号','項目','レベル','有効なカード数','カード'],krows,[6,30,30,14,40,7,13,8],wrap=(2,3,5))
NK=len(krows)+1
# --- 概要
wo=wb.create_sheet('概要',0)
wo['A1']='クリーンルーム事実票：取り込み結果の概要'; wo['A1'].font=Font(name='Arial',size=14,bold=True)
wo['A2']=f'作成：{datetime.date.today().isoformat()}（hm-cmed/cleanroom-facts の data/ から作成。件数は数式で数えています）'; wo['A2'].font=F
r0=4
wo.cell(r0,1,'■ 全体').font=FB
tot=[('カード（廃止を含む）',f'=COUNTA(カード一覧!A2:A{NC})'),('有効なカード',f'=COUNTIF(カード一覧!N2:N{NC},"<>retired")'),('廃止したカード',f'=COUNTIF(カード一覧!N2:N{NC},"retired")'),
     ('専門医の監修済み（verified_expert）',f'=COUNTIF(カード一覧!N2:N{NC},"verified_expert")'),('出典',f'=COUNTA(出典一覧!A2:A{NS})'),
     ('出題基準の項目',f'=COUNTA(出題基準の網羅!D2:D{NK})'),('カードのある項目',f'=COUNTIF(出題基準の網羅!H2:H{NK},"あり")'),('項目の網羅率',f'=B{r0+7}/B{r0+6}')]
for j,(k,v) in enumerate(tot):
    wo.cell(r0+1+j,1,k).font=F; c=wo.cell(r0+1+j,2,v); c.font=F
wo.cell(r0+8,2).number_format='0.0%'
# chapter coverage
r1=r0+11
wo.cell(r1-1,1,'■ 章ごとの出題基準の網羅').font=FB
hd=['章','章の名前','項目','レベル a','カードのある項目','網羅率','レベル a でカードのない項目']
for j,h in enumerate(hd,1):
    c=wo.cell(r1,j,h); c.font=FH; c.fill=HF
for i,ch in enumerate(chs):
    r=r1+1+i; G='出題基準の網羅!'
    vals=[ch['id'],ch['name'],f'=COUNTIF({G}$A$2:$A${NK},A{r})',f'=COUNTIFS({G}$A$2:$A${NK},A{r},{G}$F$2:$F${NK},"a")',
          f'=COUNTIFS({G}$A$2:$A${NK},A{r},{G}$H$2:$H${NK},"あり")',f'=IF(C{r}=0,0,E{r}/C{r})',f'=COUNTIFS({G}$A$2:$A${NK},A{r},{G}$F$2:$F${NK},"a",{G}$H$2:$H${NK},"なし")']
    for j,v in enumerate(vals,1): c=wo.cell(r,j,v); c.font=F
    wo.cell(r,6).number_format='0.0%'
rt=r1+1+len(chs)
wo.cell(rt,1,'計').font=FB
for j,col in [(3,'C'),(4,'D'),(5,'E'),(7,'G')]: c=wo.cell(rt,j,f'=SUM({col}{r1+1}:{col}{rt-1})'); c.font=FB
c=wo.cell(rt,6,f'=IF(C{rt}=0,0,E{rt}/C{rt})'); c.font=FB; c.number_format='0.0%'
for j in range(1,8): wo.cell(rt,j).fill=TF
# prefix table
r2=rt+3
wo.cell(r2-1,1,'■ 分野・接頭辞ごとの件数').font=FB
hd=['分野','接頭辞','範囲','有効','廃止','監修済み','OPEN','RESTRICTED','GOV_JP','GOV_US','UNKNOWN']
for j,h in enumerate(hd,1):
    c=wo.cell(r2,j,h); c.font=FH; c.fill=HF
C='カード一覧!'; rng=lambda col:f'{C}${col}$2:${col}${NC}'
present=[d for d in doms if (R/'data/facts'/f'{d[0]}.yaml').exists()]
for i,(p,n,f) in enumerate(present):
    r=r2+1+i
    vals=[f,p,n,f'=COUNTIFS({rng("C")},B{r},{rng("N")},"<>retired")',f'=COUNTIFS({rng("C")},B{r},{rng("N")},"retired")',f'=COUNTIFS({rng("C")},B{r},{rng("N")},"verified_expert")']
    vals+=[f'=COUNTIFS({rng("C")},B{r},{rng("N")},"<>retired",{rng("M")},"{k}")' for k in ['OPEN','RESTRICTED','GOV_JP','GOV_US','UNKNOWN']]
    for j,v in enumerate(vals,1): c=wo.cell(r,j,v); c.font=F
rt2=r2+1+len(present)
wo.cell(rt2,1,'計').font=FB
for j in range(4,12):
    col=get_column_letter(j); c=wo.cell(rt2,j,f'=SUM({col}{r2+1}:{col}{rt2-1})'); c.font=FB
for j in range(1,12): wo.cell(rt2,j).fill=TF
for col,w in zip('ABCDEFGHIJK',[24,30,44,8,8,9,8,11,9,9,10]): wo.column_dimensions[col].width=w
wo.freeze_panes=None
# --- 説明
wn=wb.create_sheet('説明')
notes=[('このファイルについて','hm-cmed/cleanroom-facts の data/facts（カード）、data/sources（出典）、data/curriculum（出題基準）から作った一覧です。元データを変えても、このファイルは自動では更新されません（作り直す必要があります）。'),
('件数の数え方','「概要」と「出題基準の網羅」の件数は、「カード一覧」「出典一覧」を数える数式です。'),
('有効と廃止','状態（status）が retired のカードは廃止です。置き換え先の列に、統合・差し替え先のカードがあります。ID は再利用しません。'),
('状態','draft：出典・数値の照合が未完了。verified_sources：識別子とライセンスを API で照合済み。verified_expert：専門医が監修済み（人だけが付けます）。'),
('ライセンス区分','OPEN（CC BY・CC0）、RESTRICTED（CC BY-NC・ND、出版社規約など）、GOV_JP（日本の公的資料）、GOV_US（米国連邦政府著作物）、UNKNOWN（未確認。RESTRICTED と同じ扱い）。'),
('出題基準の網羅','カードの出題基準の欄に「医学各論 番号（」が含まれ、廃止でないカードを数えています。「作らない」とした項目は、別の章の既存カードに番号を足して受けています。'),
('著作権','主張文は出典の文言を写さず自作した文です。このファイルに論文やガイドラインの本文は含みません。')]
wn.append(['項目','説明'])
for c in wn[1]: c.font=FH; c.fill=HF
for k,v in notes: wn.append([k,v])
for row in wn.iter_rows(min_row=2):
    for c in row: c.font=F; c.alignment=Alignment(wrap_text=True,vertical='top')
wn.column_dimensions['A'].width=18; wn.column_dimensions['B'].width=100
out=Path('/home/claude/cleanroom-gi-facts/reports')/f'factcards_summary_{datetime.date.today().strftime("%Y%m%d")}.xlsx'
wb.save(out); print(out, len(cards), len(S), len(krows))
